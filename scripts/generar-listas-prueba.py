"""
Genera las 20 listas de prueba de la Fase 2 (F04) en tests/fixtures/sinteticas/.

Son listas armadas por nosotros (decisión de Leandro, 12/10/2026), no listas reales: imitan
los formatos que suelen mandar los proveedores en Argentina. Cada lista va con un
`.esperado.json` que tiene las líneas que la app tiene que leer, calculadas acá a partir
de los datos de origen y no con el código de la app.

Uso: python3 scripts/generar-listas-prueba.py   (necesita openpyxl)
"""

import csv
import io
import json
import random
from pathlib import Path

from openpyxl import Workbook

DESTINO = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "sinteticas"

PRODUCTOS = [
    ("Yerba mate Playadito", "1 kg"), ("Yerba mate Taragüí con palo", "500 g"), ("Fideos Matarazzo tirabuzón", "500 g"),
    ("Fideos Lucchetti spaghetti", "500 g"), ("Arroz Gallo Oro largo fino", "1 kg"), ("Harina Pureza leudante", "1 kg"),
    ("Aceite Natura girasol", "1,5 l"), ("Aceite Cocinero mezcla", "900 ml"), ("Azúcar Ledesma", "1 kg"),
    ("Sal fina Celusal", "500 g"), ("Leche La Serenísima entera", "1 l"), ("Dulce de leche Sancor", "400 g"),
    ("Galletitas Express", "3 x 101 g"), ("Galletitas Criollitas", "3 x 100 g"), ("Café La Virginia molido", "250 g"),
    ("Té Taragüí", "x 50 saquitos"), ("Mate cocido Unión", "x 25"), ("Puré de tomate Arcor", "520 g"),
    ("Arvejas Inalpa", "350 g"), ("Atún La Campagnola al natural", "170 g"), ("Mayonesa Hellmann's", "475 g"),
    ("Lavandina Ayudín", "1 l"), ("Detergente Magistral limón", "750 ml"), ("Jabón en polvo Ala", "800 g"),
    ("Papel higiénico Higienol", "4 x 30 m"), ("Shampoo Sedal", "340 ml"), ("Desodorante Rexona", "150 ml"),
    ("Gaseosa Coca-Cola", "2,25 l"), ("Agua Villavicencio sin gas", "2 l"), ("Cerveza Quilmes lata", "473 ml"),
    ("Vino Toro tinto", "1 l"), ("Alfajor Jorgito", "x 6"), ("Chocolate Milka", "100 g"), ("Caramelos Sugus", "x 150"),
    ("Tornillo autoperforante 8x1/2", "x 100"), ("Cinta aisladora 3M", "20 m"), ("Lámpara LED 9W", "unidad"),
    ("Pila AA Duracell", "x 4"), ("Almendras peladas", "kg"), ("Nueces mariposa", "kg"), ("Avena arrollada", "kg"),
    ("Pasas de uva", "kg"), ("Semillas de chía", "kg"), ("Garbanzos", "kg"), ("Lentejas", "kg"),
]

RUBROS = ["ALMACÉN", "LIMPIEZA", "BEBIDAS", "GOLOSINAS", "PERFUMERÍA"]


def pesos_ar(centavos: int, con_signo: bool = False) -> str:
    """123450 → '1.234,50'."""
    entera = f"{centavos // 100:,}".replace(",", ".")
    texto = f"{entera},{centavos % 100:02d}"
    return f"$ {texto}" if con_signo else texto


def pesos_us(centavos: int, miles: bool = True) -> str:
    """123450 → '1,234.50' (o '1234.50' sin miles)."""
    entera = f"{centavos // 100:,}" if miles else str(centavos // 100)
    return f"{entera}.{centavos % 100:02d}"


def productos(cantidad: int, semilla: int, con_codigo: str = "A", desde: int = 1):
    """Lista de productos con código, descripción, presentación y costo en centavos (con decimales a veces)."""
    azar = random.Random(semilla)
    salida = []
    for i in range(cantidad):
        nombre, presentacion = PRODUCTOS[i % len(PRODUCTOS)]
        if i >= len(PRODUCTOS):
            nombre = f"{nombre} {['surtido', 'promo', 'familiar', 'clásico', 'light'][(i // len(PRODUCTOS)) % 5]} {i // len(PRODUCTOS)}"
        costo = azar.randint(150, 60_000) * 100 + azar.choice([0, 0, 0, 50, 25, 99])
        salida.append({"codigo": f"{con_codigo}{desde + i:04d}", "descripcion": nombre, "presentacion": presentacion, "costo": costo})
    return salida


class Lista:
    """Acumula filas del archivo y las líneas que la app tiene que leer."""

    def __init__(self):
        self.filas: list[list] = []
        self.esperadas: list[dict] = []

    def fila(self, *celdas):
        self.filas.append(list(celdas))

    def producto(self, celdas, p, codigo=True, presentacion=True, iva=None):
        self.filas.append(list(celdas))
        self.esperadas.append({
            "codigo": p["codigo"] if codigo else None,
            "descripcion": p["descripcion"],
            "presentacion": p["presentacion"] if presentacion else None,
            "costo": p["costo"],
            "ivaIncluido": iva,
            "filaOrigen": len(self.filas),
        })


def guardar_csv(nombre, lista: Lista, separador=";", codificacion="utf-8", bom=False):
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=separador, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    for fila in lista.filas:
        escritor.writerow(["" if c is None else c for c in fila])
    datos = buffer.getvalue().encode(codificacion)
    if bom:
        datos = b"\xef\xbb\xbf" + datos
    (DESTINO / nombre).write_bytes(datos)
    guardar_esperado(nombre, lista)


def guardar_xlsx(nombre, hojas: list[tuple[str, Lista]], principal: int = 0, inicio_columna: int = 1):
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, lista in hojas:
        hoja = libro.create_sheet(titulo)
        for r, fila in enumerate(lista.filas, start=1):
            for c, valor in enumerate(fila, start=inicio_columna):
                if valor is not None and valor != "":
                    hoja.cell(row=r, column=c, value=valor)
    libro.save(DESTINO / nombre)
    guardar_esperado(nombre, hojas[principal][1])


def guardar_esperado(nombre, lista: Lista):
    esperado = {"lineas": lista.esperadas}
    if len(lista.esperadas) > 500:
        # Para no llenar el repo: cantidad, suma de costos y una línea de cada 50.
        esperado = {
            "cantidad": len(lista.esperadas),
            "sumaCostos": sum(e["costo"] for e in lista.esperadas),
            "muestra": lista.esperadas[::50],
        }
    with open(DESTINO / f"{nombre}.esperado.json", "w", encoding="utf-8") as f:
        json.dump(esperado, f, ensure_ascii=False, separators=(",", ":"))
        f.write("\n")


def generar():
    DESTINO.mkdir(parents=True, exist_ok=True)
    for viejo in DESTINO.glob("*"):
        if viejo.name != "LEEME.md":
            viejo.unlink()

    # 01 · CSV simple con ; y formato argentino
    l = Lista()
    l.fila("Código", "Descripción", "Presentación", "Precio")
    for p in productos(30, 1):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], pesos_ar(p["costo"])], p)
    guardar_csv("01-almacen-basico.csv", l)

    # 02 · Excel con datos del proveedor arriba, rubros y números de Excel
    l = Lista()
    l.fila("DISTRIBUIDORA EL SOL S.A.")
    l.fila("Av. Siempreviva 742 · CUIT 30-00000000-0")
    l.fila("Lista de precios N° 41")
    l.fila("Vigencia desde 01/10/2026")
    l.fila()
    l.fila("Cód. Art.", "Descripción", "Presentación", "Precio s/IVA")
    todos = productos(40, 2)
    for i, rubro in enumerate(RUBROS[:4]):
        l.fila(None, rubro)
        for p in todos[i * 10:(i + 1) * 10]:
            l.producto([p["codigo"], p["descripcion"], p["presentacion"], p["costo"] / 100], p, iva=False)
    guardar_xlsx("02-distribuidora-con-logo.xlsx", [("Lista", l)])

    # 03 · CSV de Excel en Windows: Latin-1 y ;
    l = Lista()
    l.fila("Codigo", "Articulo", "Contenido", "P. Lista")
    for p in productos(25, 3, "B"):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], pesos_ar(p["costo"])], p)
    guardar_csv("03-windows-latin1.csv", l, codificacion="cp1252")

    # 04 · CSV con coma y punto decimal
    l = Lista()
    l.fila("SKU", "Producto", "Unidad", "Precio")
    for p in productos(30, 4, "SK"):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], pesos_us(p["costo"])], p)
    guardar_csv("04-punto-decimal.csv", l, separador=",")

    # 05 · CSV sin títulos
    l = Lista()
    for p in productos(30, 5, "1"):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], pesos_ar(p["costo"])], p)
    guardar_csv("05-sin-titulos.csv", l)

    # 06 · Excel con precio de lista y precio sugerido
    l = Lista()
    l.fila("Artículo", "Detalle", "Precio sugerido", "Precio lista", "PVP")
    for p in productos(30, 6, "C"):
        l.producto([p["codigo"], p["descripcion"], round(p["costo"] * 1.6) / 100, p["costo"] / 100, round(p["costo"] * 1.8) / 100], p, presentacion=False)
    guardar_xlsx("06-precio-lista-y-sugerido.xlsx", [("Hoja1", l)])

    # 07 · Excel con precios con IVA
    l = Lista()
    l.fila("Código", "Producto", "Precio c/IVA")
    for p in productos(30, 7, "D"):
        l.producto([p["codigo"], p["descripcion"], p["costo"] / 100], p, presentacion=False, iva=True)
    guardar_xlsx("07-precios-con-iva.xlsx", [("Precios", l)])

    # 08 · CSV exportado de un PDF: títulos repetidos en cada página y pie de página
    l = Lista()
    todos = productos(75, 8, "E")
    for pagina in range(3):
        l.fila("Código", "Descripción", "Precio")
        for p in todos[pagina * 25:(pagina + 1) * 25]:
            l.producto([p["codigo"], p["descripcion"], pesos_ar(p["costo"])], p, presentacion=False)
        l.fila(f"Página {pagina + 1} de 3")
    guardar_csv("08-titulos-repetidos.csv", l)

    # 09 · Excel con subtotales por rubro y total general
    l = Lista()
    l.fila("Código", "Descripción", "Importe")
    total_general = 0
    todos = productos(36, 9, "F")
    for i, rubro in enumerate(RUBROS[:3]):
        l.fila(None, rubro)
        subtotal = 0
        for p in todos[i * 12:(i + 1) * 12]:
            l.producto([p["codigo"], p["descripcion"], p["costo"] / 100], p, presentacion=False)
            subtotal += p["costo"]
        l.fila(None, f"Subtotal {rubro.lower()}", subtotal / 100)
        total_general += subtotal
    l.fila()
    l.fila(None, "TOTAL GENERAL", total_general / 100)
    guardar_xlsx("09-subtotales.xlsx", [("Lista", l)])

    # 10 · Texto separado por tabulaciones
    l = Lista()
    l.fila("Cod", "Descripcion", "Presentacion", "Precio unitario")
    for p in productos(30, 10, "G"):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], pesos_ar(p["costo"])], p)
    guardar_csv("10-tabulado.txt", l, separador="\t")

    # 11 · Excel con columnas en otro orden
    l = Lista()
    l.fila("Precio", "Descripción", "Código", "Envase")
    for p in productos(30, 11, "H"):
        l.producto([p["costo"] / 100, p["descripcion"], p["codigo"], p["presentacion"]], p)
    guardar_xlsx("11-columnas-desordenadas.xlsx", [("Lista", l)])

    # 12 · Excel con portada y la lista en la segunda hoja
    portada = Lista()
    portada.fila("Lista de precios octubre 2026")
    portada.fila("Pedidos al 11-5555-5555")
    l = Lista()
    l.fila("Código", "Descripción", "Unidad", "Precio")
    for p in productos(30, 12, "I"):
        l.producto([p["codigo"], p["descripcion"], p["presentacion"], p["costo"] / 100], p)
    guardar_xlsx("12-varias-hojas.xlsx", [("Portada", portada), ("Lista", l)], principal=1)

    # 13 · Excel con la tabla corrida (empieza en la columna C)
    l = Lista()
    l.fila()
    l.fila("Mayorista Norte")
    l.fila()
    l.fila("Código", "Descripción", "Precio neto")
    for p in productos(30, 13, "J"):
        l.producto([p["codigo"], p["descripcion"], p["costo"] / 100], p, presentacion=False, iva=False)
    guardar_xlsx("13-tabla-corrida.xlsx", [("Lista", l)], inicio_columna=3)

    # 14 · CSV con productos sin precio ("Consultar", "S/Stock")
    l = Lista()
    l.fila("Código", "Descripción", "Precio")
    for i, p in enumerate(productos(30, 14, "K")):
        if i % 7 == 3:
            l.fila(p["codigo"], p["descripcion"], "Consultar" if i % 2 else "S/Stock")
        else:
            l.producto([p["codigo"], p["descripcion"], pesos_ar(p["costo"])], p, presentacion=False)
    guardar_csv("14-sin-precio.csv", l)

    # 15 · Excel con precios escritos como texto con $
    l = Lista()
    l.fila("Código", "Descripción", "Precio")
    for p in productos(30, 15, "L"):
        l.producto([p["codigo"], p["descripcion"], pesos_ar(p["costo"], con_signo=True)], p, presentacion=False)
    guardar_xlsx("15-precios-como-texto.xlsx", [("Lista", l)])

    # 16 · Ferretería con muchas columnas
    l = Lista()
    l.fila("Código", "Código de barras", "Descripción", "Marca", "Unidad", "Stock", "Bonif. %", "Precio Neto", "Precio Final")
    azar = random.Random(16)
    for p in productos(30, 16, "M"):
        l.producto([p["codigo"], str(7790000000000 + azar.randint(0, 999999)), p["descripcion"], "Genérica", p["presentacion"],
                    azar.randint(0, 50), 10, p["costo"] / 100, round(p["costo"] * 1.21) / 100], p, iva=False)
    guardar_xlsx("16-ferreteria-muchas-columnas.xlsx", [("Lista", l)])

    # 17 · CSV UTF-8 con BOM, coma como separador y precios entre comillas
    l = Lista()
    l.fila("Código", "Descripción", "Precio")
    for p in productos(30, 17, "N"):
        l.producto([p["codigo"], p["descripcion"], pesos_ar(p["costo"])], p, presentacion=False)
    guardar_csv("17-utf8-bom-coma.csv", l, separador=",", bom=True)

    # 18 · Dietética: sin código, precio por kilo
    l = Lista()
    l.fila("DIETÉTICA LA HUERTA - Lista mayorista")
    l.fila()
    l.fila("Producto", "Precio x kg")
    dieteticos = [p for p in productos(45, 18) if p["presentacion"] == "kg"]
    for p in dieteticos:
        l.producto([p["descripcion"], pesos_ar(p["costo"])], p, codigo=False, presentacion=False)
    guardar_csv("18-dietetica-sin-codigo.csv", l)

    # 19 · Excel sólo con descripción y precio
    l = Lista()
    l.fila("Descripción", "Precio")
    for p in productos(30, 19):
        l.producto([p["descripcion"], p["costo"] / 100], p, codigo=False, presentacion=False)
    guardar_xlsx("19-solo-descripcion-y-precio.xlsx", [("Lista", l)])

    # 20 · Excel grande: 3.000 productos en 5 rubros
    l = Lista()
    l.fila("Lista general")
    l.fila("Código", "Descripción", "Presentación", "Precio")
    todos = productos(3000, 20, "Z")
    for i, rubro in enumerate(RUBROS):
        l.fila(None, rubro)
        for p in todos[i * 600:(i + 1) * 600]:
            l.producto([p["codigo"], p["descripcion"], p["presentacion"], p["costo"] / 100], p)
    guardar_xlsx("20-grande-3000.xlsx", [("Lista", l)])


if __name__ == "__main__":
    generar()
    print(f"Listas generadas en {DESTINO}")
